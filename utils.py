import os
import shutil
import webbrowser
import re
import sys 
from PyQt6.QtWidgets import QFileDialog, QMessageBox
from PIL import Image
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def export_to_excel_qt(headers, data, default_filename, parent):
    """Prompts user for a save location and exports data to a styled Excel file using PyQt."""
    filepath, _ = QFileDialog.getSaveFileName(
        parent,
        "Save Excel File",
        default_filename,
        "Excel Files (*.xlsx);;All Files (*)"
    )

    if not filepath:
        return

    try:
        wb = Workbook()
        ws = wb.active
        
        header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for row_data in data:
            ws.append(row_data)
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width

        wb.save(filepath)
        QMessageBox.information(parent, "Export Successful", f"Data successfully exported to\n{filepath}")

    except Exception as e:
        QMessageBox.critical(parent, "Export Failed", f"An error occurred: {e}")

def resource_path(relative_path):
    """ Get absolute path to resource, works for dev and for PyInstaller """
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

def copy_image_to_data_folder(source_path, fixed_filename=None):
    """Copies an image to the 'images' directory and returns the new path."""
    if not source_path:
        return None
    
    os.makedirs("images", exist_ok=True)
    
    filename = fixed_filename if fixed_filename else os.path.basename(source_path)
    dest_path = os.path.join("images", filename)
    
    try:
        with Image.open(source_path) as img:
            img.save(dest_path)
        return dest_path
    except Exception as e:
        print(f"Could not copy image: {e}")
        try:
            shutil.copy(source_path, dest_path)
            return dest_path
        except Exception as se:
            print(f"Shutil copy also failed: {se}")
            return None


def validate_and_format_date(date_string):
    if not date_string.strip():
        return ""
    try:
        parts = date_string.replace('-', '/').replace('.', '/').strip().split('/')
        if len(parts) != 3: return None
        day, month, year = int(parts[0]), int(parts[1]), int(parts[2])
        if year < 100: year += 2000 if year < 50 else 1900
        dt_object = datetime(year, month, day)
        return dt_object.strftime('%d/%m/%Y')
    except (ValueError, IndexError):
        return None

def generate_client_contract(client_data, company_details):
    if not client_data: return
    client_id, fname, lname, address, phone, license_type, _, company_name, company_id, bday, cin, h_prac, h_theo, v_mat, monitor, exam_date = client_data
    
    if company_details:
        comp_details_dict = {
            'name': company_details[6],
            'id': company_details[7]
        }
    else:
        comp_details_dict = {'name': 'Auto Ecole Abd El Karim', 'id': 'N/A'}

    html_content = f"""
    <!DOCTYPE html>
    <html lang="fr" dir="rtl"><head><meta charset="UTF-8"><title>Contrat de Formation - {fname} {lname}</title>
    <style>
        body {{ font-family: 'Segoe UI', sans-serif; margin: 40px; direction: rtl; text-align: right; border: 2px solid #000; padding: 30px; }}
        .header {{ text-align: center; margin-bottom: 40px; }} .header h1 {{ margin: 0; padding: 10px; border-bottom: 2px solid #000; }}
        .contract-section h2 {{ border-bottom: 1px solid #ccc; padding-bottom: 5px; }}
        .info-table {{ width: 100%; border-collapse: collapse; }} .info-table td {{ padding: 8px; border: 1px solid #ddd; }}
        .footer {{ margin-top: 50px; display: flex; justify-content: space-between; }}
        .signature-box {{ border-top: 1px solid #000; width: 250px; text-align: center; padding-top: 10px; }}
    </style></head>
    <body>
        <div class="header"><h1>عقد التكوين</h1><h2>{comp_details_dict.get('name', 'Auto Ecole Abd El Karim')}</h2></div>
        <h2>معلومات المترشح</h2>
        <table class="info-table">
            <tr><td>الإسم الكامل:</td><td>{fname} {lname}</td></tr> <tr><td>رقم البطاقة الوطنية:</td><td>{cin or 'N/A'}</td></tr>
            <tr><td>تاريخ الإزدياد:</td><td>{bday or 'N/A'}</td></tr> <tr><td>العنوان:</td><td>{address or 'N/A'}</td></tr>
        </table>
        <h2>تفاصيل التكوين</h2>
        <table class="info-table">
            <tr><td>معرف الشركة:</td><td>{company_id or 'N/A'}</td></tr>
            <tr><td>إسم الشركة:</td><td>{company_name or 'N/A'}</td></tr>
            <tr><td>نوع الرخصة:</td><td>{license_type}</td></tr>
            <tr><td>ساعات التدريب:</td><td>{h_prac} ساعة</td></tr>
            <tr><td>ساعات النظرية:</td><td>{h_theo} ساعة</td></tr>
            <tr><td>اسم المدرب:</td><td>{monitor}</td></tr>
            <tr><td>تاريخ النجاح فالإمتحان:</td><td>{exam_date}</td></tr>
        </table>
        <div class="footer"><div class="signature-box">توقيع المترشح</div><div class="signature-box">خاتم المؤسسة</div></div>
    </body></html>
    """
    contract_path = os.path.join(os.getcwd(), "images", f"contract_{client_id}.html")
    with open(contract_path, "w", encoding="utf-8") as f: f.write(html_content)
    webbrowser.open(f'file://{os.path.realpath(contract_path)}')

